"""Gnom formats and bounded signal extraction. Never log source rows."""

import base64
import csv
import hashlib
import hmac
import io
import os
import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from typing import Any

from openpyxl import load_workbook

HEADERS = (
    "Start",
    "End",
    "Title",
    "Participants",
    "Services",
    "Materials",
    "Income",
    "Outcome",
    "Address",
    "Cancelled",
    "Finished",
    "Result",
    "Created By",
    "Created",
    "Comments",
)
PESTS = {
    "клопы": r"клоп",
    "тараканы": r"таракан",
    "грызуны": r"грызун|мыш[ьи]|крыс",
    "плесень": r"плесен|плесень",
    "клещи": r"клещ",
    "муравьи": r"мурав",
}
METHODS = (
    "холодный туман",
    "горячий туман",
    "гель",
    "опрыскивание",
    "приманки",
    "ловушки",
)
CONDITIONS = ("дети", "животные", "собака", "конфиденциально", "импорт")
PHONE = re.compile(
    r"(?<!\d)(?:(?:\+?7|8)[\s().-]*)?\d{3}[\s().-]*\d{3}[\s().-]*\d{2}[\s().-]*\d{2}(?!\d)"
)


class GnomError(ValueError):
    """Only safe, fixed messages may leave the parser."""


def canonical(value: str) -> str:
    return re.sub(r"\s+", " ", value.casefold().replace("ё", "е")).strip()


def normalized_address(value: str) -> str:
    value = canonical(value)
    for pattern, replacement in [
        (r"\bг(?:ород)?\.?\s+", ""),
        (r"\bул(?:ица)?\.?\s*", "улица "),
        (r"\b(?:проспект|пр-т|пр\.)\s*", "проспект "),
        (r"\b(?:дом|д\.)\s*", "дом "),
        (r"\b(?:квартира|кв\.)\s*", "квартира "),
    ]:
        value = re.sub(pattern, replacement, value)
    return canonical(re.sub(r"[,;.]", " ", value))


def phones(value: str) -> list[str]:
    found = set()
    for match in PHONE.finditer(value):
        digits = re.sub(r"\D", "", match.group())
        if len(digits) == 11:
            digits = digits[1:]
        if len(digits) == 10:
            found.add("+7" + digits)
    return sorted(found)


def private_hash(domain: str, value: str) -> str:
    try:
        key = base64.urlsafe_b64decode(os.environ.get("PII_FERNET_KEY", ""))
    except (ValueError, TypeError):
        raise GnomError("Не настроен ключ PII") from None
    if len(key) != 32:
        raise GnomError("Не настроен ключ PII")
    derived = hmac.digest(key, ("ekodez/gnom/" + domain + "/v1").encode(), "sha256")
    return hmac.new(derived, value.encode(), "sha256").hexdigest()


def money(value: str) -> Decimal:
    value = value.strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        result = Decimal(value or "0")
        if not result.is_finite() or result < 0 or result > Decimal("999999999999.99"):
            raise InvalidOperation
        return result.quantize(Decimal("0.01"))
    except InvalidOperation:
        raise GnomError("Некорректная сумма в строке") from None


def timestamp(value: str, required: bool = False) -> datetime | None:
    if not value.strip() and not required:
        return None
    try:
        result = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if result.tzinfo:
            from zoneinfo import ZoneInfo

            result = result.astimezone(ZoneInfo("Europe/Moscow")).replace(tzinfo=None)
        return result
    except ValueError:
        for pattern in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
            try:
                return datetime.strptime(value.strip(), pattern)
            except ValueError:
                pass
    raise GnomError("Некорректная дата в строке")


def flag(value: str) -> bool:
    return canonical(value) in {"true", "1", "да", "yes", "finished", "cancelled"}


@dataclass
class GnomRow:
    start: datetime
    created: datetime | None
    end: datetime | None
    address: str
    phone_values: list[str]
    name: str
    comments: str
    income: Decimal
    outcome: Decimal
    finished: bool
    cancelled: bool
    signals: dict[str, Any]


def extract(row: dict[str, str]) -> GnomRow:
    comments = row["Comments"]
    # Ignore the entire partner line before all extractors, including phones.
    clean = re.sub(r"(?im)^.*вы\s+напарник[^\r\n]*", "", comments)
    value = canonical(clean)
    deal = re.search(r"\b(?:id|d)\s*сделки\s*[:№#]?\s*(\d+)", value)
    if not deal:
        deal = re.search(r"вы\s+взяли\s+заказ\s*[№#]?\s*(\d+)\s+в\s+работу", value)
    source = (
        "aggregator"
        if deal or "агрегаторы" in value
        else (
            "avito"
            if re.search(r"заявка\s+с\s+авито", value)
            else (
                "yandex"
                if re.search(r"заявка\s+с\s+яндекса", value)
                else "phone" if re.search(r"звонок|по телефону", value) else "other"
            )
        )
    )
    pests = [
        name
        for name, pattern in PESTS.items()
        if re.search(pattern, value + " " + canonical(row["Services"]))
    ]
    address_match = re.search(r"(?im)^\s*адрес\s*:\s*(.+)$", clean)
    address = row["Address"].strip() or (
        address_match.group(1).strip() if address_match else ""
    )
    if not address:
        raise GnomError("В строке отсутствует адрес для идентификации")
    city = next(
        (
            city
            for city in ("архангельск", "северодвинск", "новодвинск")
            if city in canonical(address)
        ),
        "",
    )
    name_match = re.search(
        r"(?im)^\s*(?:имя клиента|клиент|имя)\s*:\s*([^\r\n]+)", clean
    )
    area = re.search(r"(?<!\d)(\d+(?:[.,]\d+)?)\s*(кк|шка|ка|м[²2]|сотк\w*)\b", value)
    base = re.search(r"\bбаза\s*(\d+(?: \d{3})*(?:[.,]\d+)?)", value)
    fuel = re.search(r"\+\s*(\d+(?: \d{3})*(?:[.,]\d+)?)\s*гсм", value)
    chemical_matches = re.findall(
        r"(?im)^\s*(?:препарат|средство)\s*:\s*([^\r\n]+)", clean
    )
    signals: dict[str, Any] = {
        "source": source,
        "deal_id": deal.group(1) if deal else None,
        "pests": pests,
        "city": city,
        "methods": [method for method in METHODS if method in value],
        "conditions": [condition for condition in CONDITIONS if condition in value],
        "base_price": str(money(base.group(1))) if base else None,
        "fuel_price": str(money(fuel.group(1))) if fuel else None,
        "enhanced": "усил" in value,
        "area_value": area.group(1).replace(",", ".") if area else None,
        "area_unit": (
            (
                "rooms"
                if area.group(2) in ("кк", "шка", "ка")
                else "sotki" if area.group(2).startswith("сотк") else "m2"
            )
            if area
            else None
        ),
        "contract": bool(re.search(r"договор", value)),
        "repeat": bool(re.search(r"повторка|уже делали|закреп", value)),
        "rescheduled": "перенос" in value,
        "expense_category": (
            "ГСМ"
            if "гсм" in value
            else (
                "аренда"
                if "аренд" in value
                else (
                    "препараты"
                    if row["Materials"].strip() or chemical_matches
                    else "материалы"
                )
            )
        ),
    }
    start = timestamp(row["Start"], required=True)
    assert start is not None
    return GnomRow(
        start,
        timestamp(row["Created"]),
        timestamp(row["End"]),
        address,
        phones(
            re.sub(
                r"(?im)^.*(?:\b(?:id|d)\s*сделки|вы\s+взяли\s+заказ)[^\r\n]*", "", clean
            )
        ),
        name_match.group(1).strip() if name_match else "",
        comments,
        money(row["Income"]),
        money(row["Outcome"]),
        flag(row["Finished"]),
        flag(row["Cancelled"]),
        {
            **signals,
            "chemical_candidates": chemical_matches
            + ([row["Materials"].strip()] if row["Materials"].strip() else []),
        },
    )


class HtmlRows(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[str]] = []
        self.row: list[str] = []
        self.cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "tr":
            self.row = []
        if tag in ("td", "th"):
            self.cell = []
        if tag == "br" and self.cell is not None:
            self.cell.append("\n")

    def handle_data(self, data: str) -> None:
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self.cell is not None:
            self.row.append("".join(self.cell))
            self.cell = None
        if tag == "tr" and self.row:
            self.rows.append(self.row)


def parse_gnom(content: bytes) -> list[GnomRow]:
    if len(content) > 20 * 1024 * 1024:
        raise GnomError("Файл превышает 20 МБ")
    try:
        if content.startswith(b"PK"):
            import zipfile

            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if (
                    sum(item.file_size for item in archive.infolist())
                    > 100 * 1024 * 1024
                ):
                    raise GnomError("Распакованный файл превышает лимит")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                rows = (
                    [
                        ["" if value is None else str(value) for value in row]
                        for row in book.active.iter_rows(values_only=True)
                    ]
                    if book.active
                    else []
                )
            finally:
                book.close()
        else:
            try:
                decoded = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                decoded = content.decode("cp1251")
            if re.search(r"<table\b", decoded, re.I):
                parser = HtmlRows()
                parser.feed(decoded)
                rows = parser.rows
            else:
                dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=",;\t")
                rows = list(csv.reader(io.StringIO(decoded), dialect))
        expected = {canonical(header) for header in HEADERS}
        for index, cells in enumerate(rows[:40]):
            normalized = [canonical(cell) for cell in cells]
            if expected.issubset(normalized):
                if len(set(normalized)) != len(normalized):
                    raise GnomError("Повторяющиеся заголовки")
                if len(rows) - index > 50001:
                    raise GnomError("Слишком много строк")
                result = []
                for values in rows[index + 1 :]:
                    if not any(value.strip() for value in values):
                        continue
                    mapped = {
                        header: (
                            values[normalized.index(canonical(header))]
                            if normalized.index(canonical(header)) < len(values)
                            else ""
                        )
                        for header in HEADERS
                    }
                    result.append(extract(mapped))
                return result
        raise GnomError("Не найдены обязательные заголовки Gnom")
    except GnomError:
        raise
    except Exception:
        raise GnomError("Не удалось прочитать формат Gnom") from None


def signature(row: GnomRow) -> tuple[str, str, str]:
    phone_hash = private_hash("phone", row.phone_values[0]) if row.phone_values else ""
    address_hash = private_hash("address", normalized_address(row.address))
    digest = hashlib.sha256(
        f"{row.start.isoformat()}|{phone_hash}|{address_hash}".encode()
    ).hexdigest()
    return digest, phone_hash, address_hash
