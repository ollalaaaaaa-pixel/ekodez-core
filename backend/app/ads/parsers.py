import csv
import re
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

from app.ads.config import ALLOWED_PLATFORMS
from app.security.ad_phone_hash import PhoneHashError, phone_hmac


class AdsParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedSpendRow:
    campaign: str
    period_start: date
    period_end: date
    spend: Decimal
    impressions: int
    clicks: int
    conversions: int


@dataclass(frozen=True)
class ParsedCallRow:
    call_date: datetime
    phone_hash: str


@dataclass(frozen=True)
class ParsedAdsFile:
    spend_rows: tuple[ParsedSpendRow, ...]
    call_rows: tuple[ParsedCallRow, ...]


ALIASES = {
    "campaign": {"кампания", "название кампании", "campaign"},
    "start": {"начало периода", "дата начала", "period start"},
    "end": {"конец периода", "дата окончания", "period end"},
    "date": {"дата", "день", "date"},
    "spend": {"расход", "расход руб", "затраты", "spend", "cost"},
    "impressions": {"показы", "impressions"},
    "clicks": {"клики", "переходы", "clicks"},
    "conversions": {"конверсии", "целевые действия", "заявки", "conversions"},
    "call_date": {"дата звонка", "время звонка", "call date"},
    "phone": {"телефон", "номер телефона", "phone"},
}


def _header(value: object) -> str:
    text = str(value or "").strip().lower().replace("ё", "е")
    text = re.sub(r"[^\w]+", " ", text, flags=re.UNICODE)
    return " ".join(text.split())


def _field_map(headers: list[object]) -> dict[str, int]:
    normalized = [_header(value) for value in headers]
    result: dict[str, int] = {}
    for field, aliases in ALIASES.items():
        for index, value in enumerate(normalized):
            if value in aliases:
                result[field] = index
                break
    return result


def _decimal(value: object) -> Decimal:
    if isinstance(value, Decimal):
        return value.quantize(Decimal("0.01"))
    text = str(value or "").strip().replace("\xa0", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(",", "")
    else:
        text = text.replace(",", ".")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation as exc:
        raise AdsParseError("invalid spend value") from exc


def _integer(value: object) -> int:
    if value in (None, ""):
        return 0
    try:
        return int(Decimal(str(value).replace(" ", "").replace(",", ".")))
    except (InvalidOperation, ValueError) as exc:
        raise AdsParseError("invalid counter value") from exc


def _date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for pattern in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    raise AdsParseError("invalid date")


def _datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time())
    text = str(value or "").strip()
    for pattern in (
        "%d.%m.%Y %H:%M:%S",
        "%d.%m.%Y %H:%M",
        "%Y-%m-%d %H:%M:%S",
        "%d.%m.%Y",
        "%Y-%m-%d",
    ):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    raise AdsParseError("invalid call date")


class _TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.rows: list[list[str]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() == "tr":
            self._row = []
        elif tag.lower() in {"td", "th"} and self._row is not None:
            self._cell = []

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"td", "th"} and self._cell is not None:
            assert self._row is not None
            self._row.append("".join(self._cell).strip())
            self._cell = None
        elif tag.lower() == "tr" and self._row is not None:
            if self._row:
                self.rows.append(self._row)
            self._row = None


def _read_rows(path: Path) -> list[list[object]]:
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            if sheet is None:
                return []
            return [list(row) for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    if suffix == ".xls":
        parser = _TableParser()
        try:
            parser.feed(path.read_text(encoding="utf-8-sig"))
        except (OSError, UnicodeDecodeError) as exc:
            raise AdsParseError("legacy XLS must contain an HTML table") from exc
        return [list(row) for row in parser.rows]
    if suffix == ".csv":
        try:
            text = path.read_text(encoding="utf-8-sig")
        except (OSError, UnicodeDecodeError) as exc:
            raise AdsParseError("cannot read CSV") from exc
        if not text.strip():
            return []
        if ";" in text:
            delimiter = ";"
        elif "\t" in text:
            delimiter = "\t"
        else:
            try:
                delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",").delimiter
            except csv.Error:
                delimiter = ","
        return [list(row) for row in csv.reader(text.splitlines(), delimiter=delimiter)]
    raise AdsParseError("unsupported advertising file format")


def _value(row: list[object], fields: dict[str, int], name: str) -> object:
    index = fields[name]
    return row[index] if index < len(row) else None


def _nonempty_rows(rows: Iterable[list[object]]) -> list[list[object]]:
    return [row for row in rows if any(str(value or "").strip() for value in row)]


def parse_ads_file(
    path: Path, platform: str, pii_key: str | None = None
) -> ParsedAdsFile:
    if platform not in ALLOWED_PLATFORMS:
        raise AdsParseError("unknown advertising platform")
    rows = _nonempty_rows(_read_rows(path))
    if len(rows) < 2:
        raise AdsParseError("advertising file is empty")
    header_index = -1
    fields: dict[str, int] = {}
    spend_shape = False
    call_shape = False
    for index, candidate in enumerate(rows[:40]):
        candidate_fields = _field_map(candidate)
        candidate_spend = "spend" in candidate_fields and (
            "date" in candidate_fields or "start" in candidate_fields
        )
        candidate_call = {"call_date", "phone"} <= candidate_fields.keys()
        if candidate_spend or candidate_call:
            header_index = index
            fields = candidate_fields
            spend_shape = candidate_spend
            call_shape = candidate_call
            break
    if not spend_shape and not call_shape:
        raise AdsParseError("advertising columns were not recognized")

    spend_rows: list[ParsedSpendRow] = []
    call_rows: list[ParsedCallRow] = []
    for row in rows[header_index + 1 :]:
        if spend_shape:
            start = _date(_value(row, fields, "start" if "start" in fields else "date"))
            end = _date(_value(row, fields, "end")) if "end" in fields else start
            spend_rows.append(
                ParsedSpendRow(
                    campaign=(
                        str(_value(row, fields, "campaign")).strip()
                        if "campaign" in fields
                        else platform
                    ),
                    period_start=start,
                    period_end=end,
                    spend=_decimal(_value(row, fields, "spend")),
                    impressions=(
                        _integer(_value(row, fields, "impressions"))
                        if "impressions" in fields
                        else 0
                    ),
                    clicks=(
                        _integer(_value(row, fields, "clicks"))
                        if "clicks" in fields
                        else 0
                    ),
                    conversions=(
                        _integer(_value(row, fields, "conversions"))
                        if "conversions" in fields
                        else 0
                    ),
                )
            )
        if call_shape:
            try:
                digest = phone_hmac(str(_value(row, fields, "phone")), key=pii_key)
            except PhoneHashError as exc:
                raise AdsParseError("call phone cannot be protected") from exc
            call_rows.append(
                ParsedCallRow(
                    call_date=_datetime(_value(row, fields, "call_date")),
                    phone_hash=digest,
                )
            )
    return ParsedAdsFile(tuple(spend_rows), tuple(call_rows))
