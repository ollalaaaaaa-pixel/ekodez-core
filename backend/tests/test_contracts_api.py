import json
import os
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch

from cryptography.fernet import Fernet
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, select
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app import main
from app.models import Base, Client, ContractPeriod, Transaction
from app.security.pii import encrypt_sensitive_mapping


class ContractsAndActsApiTest(unittest.TestCase):
    def _new_package(self):
        contract = self._create_contract()
        path = f"/api/contracts/{contract['id']}/package/2026-09"
        empty = self.client.get(path).json()
        self.assertIsNone(empty["period"])
        self.assertIsNone(empty["inspection"])
        self.assertEqual(self.client.get(path).json(), empty)
        response = self.client.patch(
            path,
            json={
                "expected_revision": empty["revision"],
                "inspection": {
                    "inspection_date": "2026-09-25",
                    "control_date": "2026-09-26",
                    "ksp_count": 9,
                },
                "period": {
                    "invoice_date": "2026-09-30",
                    "preparations": "ТЕСТ средство",
                    "extra_services": ["Услуга, с запятой"],
                },
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        return path, response.json()

    def test_package_load_defaults_and_partial_update(self):
        path, saved = self._new_package()
        self.assertEqual(saved, self.client.get(path).json())
        self.assertEqual(saved["inspection"]["ksp_count"], 9)
        self.assertEqual(saved["period"]["infestation_degree"], "начальная")
        self.assertEqual(saved["inspection"]["rodents_caught"], 0)
        changed = self.client.patch(
            path,
            json={
                "expected_revision": saved["revision"],
                "period": {"infestation_degree": "Не обнаружено"},
            },
        )
        self.assertEqual(changed.status_code, 200, changed.text)
        data = changed.json()
        self.assertEqual(data["inspection"], saved["inspection"])
        for field in saved["period"]:
            if field != "infestation_degree":
                self.assertEqual(data["period"][field], saved["period"][field], field)
        self.assertEqual(
            self.client.patch(
                path, json={"expected_revision": data["revision"]}
            ).json(),
            data,
        )
        self.assertEqual(
            self.client.patch(
                path,
                json={
                    "expected_revision": saved["revision"],
                    "period": {"preparations": "stale"},
                },
            ).status_code,
            409,
        )

    def test_protected_package_confirmation_versions_and_noop(self):
        path, saved = self._new_package()
        self._save_billing_client()
        with (
            tempfile.TemporaryDirectory() as temp_dir,
            patch.object(main, "DOCUMENT_OUTPUT_ROOT", Path(temp_dir)),
            patch.object(
                main,
                "_document_profile",
                return_value={
                    "EXECUTOR_BANK_DETAILS": "ТЕСТ",
                    "EXECUTOR_INN": "ТЕСТ",
                    "EXECUTOR_OGRNIP": "ТЕСТ",
                    "TAX_MODE": "НДС не облагается (УСН)",
                },
            ),
        ):
            signed = self.client.patch(
                path,
                json={
                    "expected_revision": saved["revision"],
                    "generate": True,
                    "inspection": {
                        "status": "signed",
                        "signed_at": "2026-09-30T12:00:00",
                    },
                    "period": {
                        "work_act_status": "signed",
                        "work_act_signed_at": "2026-09-30T12:00:00",
                    },
                },
            )
            self.assertEqual(signed.status_code, 200, signed.text)
            original = signed.json()
            blocked = self.client.patch(
                f"/api/contract-periods/{original['period']['id']}",
                json={"preparations": "ТЕСТ обход"},
            )
            self.assertEqual(blocked.status_code, 409)
            self.assertEqual(original["period"]["file_manifest"][0]["version"], 1)
            old_files = {str(p): p.read_bytes() for p in Path(temp_dir).rglob("*.docx")}
            noop = self.client.patch(
                path, json={"expected_revision": original["revision"], "generate": True}
            )
            self.assertEqual(noop.json(), original)
            payload = {
                "expected_revision": original["revision"],
                "period": {"preparations": "ТЕСТ исправлено"},
            }
            denied = self.client.patch(path, json=payload)
            self.assertEqual(denied.status_code, 409)
            self.assertEqual(self.client.get(path).json(), original)
            approved = self.client.patch(path, json={**payload, "confirm_edit": True})
            self.assertEqual(approved.status_code, 200, approved.text)
            updated = approved.json()
            self.assertEqual(updated["period"]["file_manifest"][0]["version"], 2)
            self.assertEqual(updated["inspection"], original["inspection"])
            self.assertEqual(
                updated["period"]["work_act_signed_at"],
                original["period"]["work_act_signed_at"],
            )
            for name, content in old_files.items():
                self.assertEqual(Path(name).read_bytes(), content)
            with patch.object(
                main,
                "build_month_package",
                side_effect=main.DocumentTemplateError("failed"),
            ):
                failed = self.client.patch(
                    path,
                    json={
                        "expected_revision": updated["revision"],
                        "confirm_edit": True,
                        "inspection": {"ksp_count": 1},
                    },
                )
                self.assertEqual(failed.status_code, 422)
                self.assertEqual(self.client.get(path).json(), updated)

    def test_package_edit_localhost_only(self):
        path, saved = self._new_package()
        remote = TestClient(main.app, client=("192.168.1.20", 51000))
        try:
            response = remote.patch(
                path,
                json={"expected_revision": saved["revision"]},
                headers={"X-Forwarded-For": "127.0.0.1"},
            )
        finally:
            remote.close()
        self.assertEqual(response.status_code, 403)

    def test_invoice_money_words(self):
        cases = {
            "0.00": "ноль рублей 00 копеек",
            "1.01": "один рубль 01 копейка",
            "2.02": "два рубля 02 копейки",
            "11.11": "одиннадцать рублей 11 копеек",
            "3000.00": "три тысячи рублей 00 копеек",
            "21001.21": "двадцать одна тысяча один рубль 21 копейка",
            "1500000.50": "один миллион пятьсот тысяч рублей 50 копеек",
            "1000000000.00": "один миллиард рублей 00 копеек",
        }
        for amount, expected in cases.items():
            with self.subTest(amount=amount):
                self.assertEqual(main._money_words(Decimal(amount)), expected)
        for amount in ("-1.00", "1000000000000.00"):
            with self.assertRaises(main.DocumentTemplateError):
                main._money_words(Decimal(amount))

    def setUp(self):
        self.original_engine = main.engine
        self.engine = create_engine(
            "sqlite:///:memory:",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        main.engine = self.engine
        self.key = Fernet.generate_key().decode("ascii")
        self.environment = patch.dict(
            os.environ, {"PII_FERNET_KEY": self.key}, clear=False
        )
        self.environment.start()
        self.client = TestClient(main.app, client=("127.0.0.1", 51000))
        self.object_id = self.client.post(
            "/api/objects",
            json={
                "name": "ТЕСТ Хостел",
                "address": "г. Архангельск, пр. Тестовый, 1",
                "type": "other",
                "area_sqm": "90.00",
                "contract": None,
                "risk_points": [],
                "status": "active",
            },
        ).json()["id"]

    def tearDown(self):
        self.client.close()
        self.environment.stop()
        main.engine = self.original_engine
        self.engine.dispose()

    def _create_contract(self) -> dict[str, object]:
        response = self.client.patch(
            f"/api/objects/{self.object_id}",
            json={
                "contract": {
                    "number": "ТЕСТ-01/09/26",
                    "contract_date": "2026-09-01",
                    "price": "5000.00",
                    "inspection_price": "3000.00",
                    "periodicity": "semiannual",
                    "service_months": [3, 9],
                }
            },
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response.json()["contract"]

    def _save_billing_client(self) -> None:
        response = self.client.put(
            f"/api/objects/{self.object_id}/billing-client",
            json={
                "client_type": "legal_entity",
                "name": "ООО «ТЕСТ Хостел»",
                "phone": "+7 921 111-22-33",
                "representative": "Кузнецова Ольга Викторовна",
                "representative_role": "директор",
                "inn": "2901000010",
                "kpp": "290101001",
                "registration_number": "1022900000000",
                "legal_address": "г. Архангельск, пр. Тестовый, 1",
                "bank_details": "р/с 40702810000000000000",
            },
        )
        self.assertEqual(response.status_code, 200, response.text)

    def test_contract_requires_manual_price_and_owner_selected_months(self):
        missing_price = self.client.patch(
            f"/api/objects/{self.object_id}",
            json={
                "contract": {
                    "number": "ТЕСТ-без-цены",
                    "periodicity": "monthly",
                }
            },
        )
        self.assertEqual(missing_price.status_code, 422)

        missing_months = self.client.patch(
            f"/api/objects/{self.object_id}",
            json={
                "contract": {
                    "number": "ТЕСТ-без-месяцев",
                    "price": "5000.00",
                    "periodicity": "semiannual",
                    "service_months": [],
                }
            },
        )
        self.assertEqual(missing_months.status_code, 422)

        contract = self._create_contract()
        self.assertEqual(contract["price"], "5000.00")
        self.assertEqual(contract["inspection_price"], "3000.00")
        self.assertEqual(contract["service_months"], [3, 9])
        self.assertNotIn("monthly_amount", contract)

    def test_client_representation_uses_document_wording(self):
        self.assertEqual(
            main._client_representation(
                {
                    "client_type": "sole_proprietor",
                    "name": "ИП Шумилова Инга Владимировна",
                }
            ),
            "ИП Шумилова Инга Владимировна",
        )
        self.assertEqual(
            main._client_representation(
                {
                    "client_type": "legal_entity",
                    "name": "ООО «ТЕСТ Хостел»",
                    "representative": "Кузнецова Ольга Викторовна",
                    "representative_role": "генеральный директор",
                }
            ),
            "ООО «ТЕСТ Хостел» в лице генерального директора Кузнецовой О.В.",
        )

    def test_contract_money_edits_from_lan_are_forbidden_and_proxy_is_ignored(self):
        contract = self._create_contract()
        payload = {key: value for key, value in contract.items() if key != "id"}
        payload["price"] = "6000.00"
        payload["inspection_price"] = "3500.00"

        remote = TestClient(main.app, client=("192.168.1.20", 51000))
        response = remote.patch(
            f"/api/objects/{self.object_id}",
            json={"contract": payload},
            headers={"X-Forwarded-For": "127.0.0.1"},
        )
        remote.close()

        self.assertEqual(response.status_code, 403)
        self.assertEqual(
            response.json()["detail"],
            "Изменение денежных полей договора доступно только "
            "на компьютере владельца",
        )
        stored = self.client.get(f"/api/objects/{self.object_id}").json()["contract"]
        self.assertEqual(stored["price"], "5000.00")
        self.assertEqual(stored["inspection_price"], "3000.00")

    def test_contract_money_edits_from_localhost_are_allowed(self):
        contract = self._create_contract()
        payload = {key: value for key, value in contract.items() if key != "id"}
        payload["price"] = "6000.00"
        payload["inspection_price"] = "3500.00"

        response = self.client.patch(
            f"/api/objects/{self.object_id}", json={"contract": payload}
        )

        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["contract"]["price"], "6000.00")
        self.assertEqual(response.json()["contract"]["inspection_price"], "3500.00")

    def test_billing_requisites_are_masked_and_local_reveal_is_audited(self):
        payload = {
            "client_type": "legal_entity",
            "name": "ООО «ТЕСТ Хостел»",
            "phone": "+7 921 111-22-33",
            "representative": "Кузнецова Ольга Викторовна",
            "representative_role": "директор",
            "inn": "2901000010",
            "kpp": "290101001",
            "registration_number": "1022900000000",
            "legal_address": "г. Архангельск, пр. Тестовый, 1",
            "bank_details": "р/с 40702810000000000000",
        }
        saved = self.client.put(
            f"/api/objects/{self.object_id}/billing-client", json=payload
        )
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertNotEqual(saved.json()["inn"], payload["inn"])
        self.assertNotIn("407028", saved.json()["bank_details"])

        remote = TestClient(main.app, client=("192.168.1.20", 51000))
        denied = remote.get(
            f"/api/objects/{self.object_id}/billing-client?show_pii=true"
        )
        remote.close()
        self.assertEqual(denied.status_code, 403)

        revealed = self.client.get(
            f"/api/objects/{self.object_id}/billing-client?show_pii=true"
        )
        self.assertEqual(revealed.status_code, 200)
        self.assertEqual(revealed.json()["inn"], payload["inn"])
        self.assertEqual(revealed.headers["cache-control"], "no-store")
        with Session(self.engine) as session:
            row = session.scalar(select(Client))
            assert row is not None
            self.assertNotIn(payload["inn"], row.inn_masked or "")
            self.assertNotIn("407028", row.bank_details_masked or "")
            self.assertIsNotNone(row.encrypted_requisites)

    def test_inspection_period_defaults_updates_and_explicit_income_link(self):
        contract = self._create_contract()
        report = self.client.post(
            f"/api/contracts/{contract['id']}/inspection-reports/2026-09",
            json={"control_date": "2026-09-25"},
        )
        self.assertEqual(report.status_code, 200, report.text)
        self.assertEqual(report.json()["control_date"], "2026-09-25")
        self.assertEqual(report.json()["ksp_count"], 5)
        self.assertEqual(report.json()["disinsection_glue_count"], 6)

        period = self.client.post(
            f"/api/contracts/{contract['id']}/periods/2026-09", json={}
        )
        self.assertEqual(period.status_code, 200, period.text)
        self.assertTrue(period.json()["paid_service_due"])
        self.assertEqual(period.json()["price_snapshot"], "5000.00")
        self.assertEqual(period.json()["infestation_degree"], "начальная")
        self.assertEqual(period.json()["extra_services"], [])
        self.assertEqual(period.json()["invoice_number"], "1")

        edited = self.client.patch(
            f"/api/contract-periods/{period.json()['id']}",
            json={
                "preparations": "ТЕСТ препарат",
                "infestation_degree": "средняя",
                "extra_services": ["ТЕСТ дополнительный осмотр"],
                "invoice_number": "77-А",
            },
        )
        self.assertEqual(edited.status_code, 200, edited.text)
        self.assertEqual(edited.json()["invoice_number"], "77-А")
        signed = self.client.patch(
            f"/api/contract-periods/{period.json()['id']}",
            json={
                "work_act_status": "signed",
                "work_act_signed_at": "2026-09-30T12:00:00+03:00",
            },
        )
        self.assertEqual(signed.status_code, 200, signed.text)
        timeline = self.client.get(
            f"/api/objects/{self.object_id}/contract-timeline"
        ).json()
        self.assertIn("work_act_signed", {item["type"] for item in timeline})

        with Session(self.engine) as session:
            session.add(
                Transaction(
                    source="manual",
                    operation_date=date(2026, 9, 30),
                    amount=Decimal("5000.00"),
                    kind="income",
                    review_required=False,
                    needs_review=False,
                    object_id=self.object_id,
                )
            )
            session.commit()
            transaction_id = session.scalar(select(Transaction.id))

        linked = self.client.patch(
            f"/api/contract-periods/{period.json()['id']}",
            json={"transaction_id": transaction_id},
        )
        self.assertEqual(linked.status_code, 200, linked.text)
        self.assertEqual(linked.json()["transaction_id"], transaction_id)

        with Session(self.engine) as session:
            stored = session.get(ContractPeriod, period.json()["id"])
            assert stored is not None
            self.assertEqual(stored.transaction_id, transaction_id)

    def test_package_generation_is_localhost_only_and_records_manifest(self):
        contract = self._create_contract()
        self._save_billing_client()
        report = self.client.post(
            f"/api/contracts/{contract['id']}/inspection-reports/2026-09",
            json={
                "inspection_date": "2026-09-25",
                "control_date": "2026-09-26",
            },
        )
        self.assertEqual(report.status_code, 200, report.text)
        period = self.client.post(
            f"/api/contracts/{contract['id']}/periods/2026-09",
            json={"invoice_date": "2026-09-25", "preparations": "ТЕСТ препарат"},
        ).json()

        with tempfile.TemporaryDirectory() as temp_dir:
            profile_path = Path(temp_dir) / "company-profile.json"
            profile_values = {
                "EXECUTOR_BANK_DETAILS": "ТЕСТ банк",
                "EXECUTOR_INN": "ТЕСТ ИНН",
                "EXECUTOR_OGRNIP": "ТЕСТ ОГРНИП",
                "TAX_MODE": "НДС не облагается (УСН)",
            }
            encrypted_profile = encrypt_sensitive_mapping(profile_values)
            self.assertIsNotNone(encrypted_profile)
            profile_path.write_text(
                json.dumps({"encrypted_profile": encrypted_profile}),
                encoding="utf-8",
            )
            profile_text = profile_path.read_text(encoding="utf-8")
            self.assertNotIn("ТЕСТ ИНН", profile_text)
            self.assertNotIn("ТЕСТ банк", profile_text)
            with (
                patch.object(main, "DOCUMENT_OUTPUT_ROOT", Path(temp_dir) / "out"),
                patch.object(main, "DOCUMENT_PROFILE_PATH", profile_path),
            ):
                remote = TestClient(main.app, client=("192.168.1.20", 51000))
                denied = remote.post(f"/api/contract-periods/{period['id']}/generate")
                remote.close()
                generated = self.client.post(
                    f"/api/contract-periods/{period['id']}/generate"
                )

            self.assertEqual(denied.status_code, 403)
            self.assertEqual(generated.status_code, 200, generated.text)
            self.assertEqual(len(generated.json()["file_manifest"]), 3)
            self.assertEqual(generated.json()["file_manifest"][0]["version"], 1)
            self.assertNotIn(str(Path(temp_dir)), generated.text)
            package_dir = Path(temp_dir) / "out" / "2026-09" / "ТЕСТ Хостел" / "v1"
            package_text = "\n".join(
                paragraph.text
                for file_path in package_dir.glob("*.docx")
                for paragraph in Document(str(file_path)).paragraphs
            )
            self.assertIn("НДС не облагается (УСН)", package_text)
            self.assertNotIn("БЕЗ НДС", package_text)
            self.assertNotIn("Претензий нет..", package_text)
            invoice = Document(str(package_dir / "Счёт.docx"))
            invoice_text = "\n".join(
                [p.text for p in invoice.paragraphs]
                + [c.text for t in invoice.tables for r in t.rows for c in r.cells]
            )
            self.assertNotIn("legal_entity", invoice_text)
            self.assertNotIn("sole_proprietor", invoice_text)
            self.assertNotIn("рублей рублей", invoice_text)
            self.assertIn("пять тысяч рублей 00 копеек", invoice_text)
            for file_name in ("Акт_выполненных_работ.docx", "Счёт.docx"):
                document = Document(str(package_dir / file_name))
                total_row = next(
                    row
                    for table in document.tables
                    for row in table.rows
                    if row.cells[0].text.strip().startswith("Итого")
                )
                self.assertEqual(total_row.cells[0].text.strip(), "Итого")
                self.assertTrue(
                    all(not cell.text.strip() for cell in total_row.cells[1:-1])
                )
                self.assertEqual(total_row.cells[-1].text.strip(), "5000.00")
            first_name = generated.json()["file_manifest"][0]["name"]
            with (
                patch.object(main, "DOCUMENT_OUTPUT_ROOT", Path(temp_dir) / "out"),
                patch.object(main, "DOCUMENT_PROFILE_PATH", profile_path),
            ):
                downloaded = self.client.get(
                    f"/api/contract-periods/{period['id']}/files/{first_name}"
                )
                remote = TestClient(main.app, client=("192.168.1.20", 51000))
                remote_download = remote.get(
                    f"/api/contract-periods/{period['id']}/files/{first_name}"
                )
                remote.close()
            self.assertEqual(downloaded.status_code, 200)
            self.assertGreater(len(downloaded.content), 0)
            self.assertEqual(remote_download.status_code, 403)
            with (
                patch.object(main, "DOCUMENT_OUTPUT_ROOT", Path(temp_dir) / "out"),
                patch.object(main, "DOCUMENT_PROFILE_PATH", profile_path),
            ):
                regenerated = self.client.post(
                    f"/api/contract-periods/{period['id']}/generate"
                )
            self.assertEqual(regenerated.status_code, 200, regenerated.text)
            self.assertEqual(regenerated.json()["file_manifest"][0]["version"], 2)

    def test_document_profile_is_encrypted_and_localhost_only(self):
        payload = {
            "executor_bank_details": "ТЕСТ банк и счёт",
            "executor_inn": "ТЕСТ ИНН",
            "executor_ogrnip": "ТЕСТ ОГРНИП",
            "tax_mode": "НДС не облагается",
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            profile_path = Path(temp_dir) / "company-profile.json"
            with patch.object(main, "DOCUMENT_PROFILE_PATH", profile_path):
                remote = TestClient(main.app, client=("192.168.1.20", 51000))
                denied = remote.put("/api/document-profile", json=payload)
                remote.close()
                saved = self.client.put("/api/document-profile", json=payload)
                status = self.client.get("/api/document-profile/status")
            self.assertEqual(denied.status_code, 403)
            self.assertEqual(saved.json(), {"status": "configured"})
            self.assertEqual(status.json(), {"configured": True})
            stored = profile_path.read_text(encoding="utf-8")
            self.assertNotIn("ТЕСТ банк", stored)
            self.assertNotIn("ТЕСТ ИНН", stored)

    def test_auto_packages_trigger_is_local_and_idempotent(self):
        self._create_contract()
        self._save_billing_client()
        with tempfile.TemporaryDirectory() as temp_dir:
            profile_path = Path(temp_dir) / "company-profile.json"
            encrypted_profile = encrypt_sensitive_mapping(
                {
                    "EXECUTOR_BANK_DETAILS": "ТЕСТ банк",
                    "EXECUTOR_INN": "ТЕСТ ИНН",
                    "EXECUTOR_OGRNIP": "ТЕСТ ОГРНИП",
                    "TAX_MODE": "НДС не облагается",
                }
            )
            assert encrypted_profile is not None
            profile_path.write_text(
                json.dumps({"encrypted_profile": encrypted_profile}), encoding="utf-8"
            )
            with (
                patch.object(main, "DOCUMENT_OUTPUT_ROOT", Path(temp_dir) / "out"),
                patch.object(main, "DOCUMENT_PROFILE_PATH", profile_path),
            ):
                remote = TestClient(main.app, client=("192.168.1.20", 51000))
                denied = remote.post("/api/contracts/auto-packages?month=2026-09")
                remote.close()
                first = self.client.post("/api/contracts/auto-packages?month=2026-09")
                second = self.client.post("/api/contracts/auto-packages?month=2026-09")

            self.assertEqual(denied.status_code, 403)
            self.assertEqual(first.status_code, 200, first.text)
            self.assertEqual(first.json()["ready"], ["ТЕСТ Хостел"])
            self.assertEqual(second.json()["ready"], [])
            with Session(self.engine) as session:
                period = session.scalar(select(ContractPeriod))
                assert period is not None
                self.assertEqual(period.price_snapshot, Decimal("3000.00"))
                self.assertEqual(len(period.file_manifest), 2)


if __name__ == "__main__":
    unittest.main()
